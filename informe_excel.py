"""
informe_excel.py — Lectura del Excel de entrada, informe Excel de cada ejecución y plantilla.

ENTRADA: leer_excel()
---------------------
Lee un fichero .xlsx y devuelve sus filas como diccionarios.
  - Hoja: la llamada "Usuarios" (la de la plantilla) o "Resultados" (la del informe
    de cambiar_correos.py, para encadenar con reiniciar_password.py). Si no hay
    ninguna de las dos, se usa la primera hoja.
  - Cabecera: la primera fila con contenido. Los nombres de columna se normalizan
    a minúsculas, sin tildes y con "_" en vez de espacios: "Correo antiguo"
    pasa a ser "correo_antiguo".
  - Se saltan las filas vacías. Cada fila lleva además "_fila", su número de fila
    real en Excel, para poder localizarla en el informe.

INFORME: guardar_informe()
--------------------------
Crea un Excel con cuatro hojas:
  Resumen       Datos de la ejecución (quién, cuándo, modo, duración...) y totales.
  Resultados    Una fila por cada fila de entrada procesada, coloreada según su estado.
  Registro API  Cada petición hecha a la API, en orden: acción, ruta, código HTTP,
                intentos, duración y resultado. Nunca incluye contraseñas.
  Leyenda       Qué significa cada estado, cada caso y cada columna.

PLANTILLA: crear_plantilla()
----------------------------
Genera el Excel de entrada vacío (hoja "Usuarios" con desplegable de rol y hoja
"Instrucciones"). Desde la terminal:  python3 informe_excel.py plantilla_entrada.xlsx
"""
from __future__ import annotations

import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = "1F4E78"
# Color de fondo de cada fila de la hoja Resultados según su estado.
COLORES_ESTADO = {
    "OK": "C6EFCE",             # verde
    "SIMULADO": "DDEBF7",       # azul
    "OMITIDO": "EDEDED",        # gris
    "PENDIENTE": "FFEB9C",      # amarillo
    "NO ENCONTRADO": "FFEB9C",  # amarillo
    "ERROR": "FFC7CE",          # rojo
    "NO PROCESADO": "F8CBAD",   # naranja
}
COLUMNAS_PLANTILLA = ["Correo antiguo", "Correo nuevo", "Nombre", "Apellidos", "Rol"]
ROLES_PLANTILLA = ["alumno", "profesor", "admin colegio"]


@dataclass
class Columna:
    """Definición de una columna del informe: clave en los datos, título, ancho y
    descripción (la descripción aparece en la hoja Leyenda)."""
    clave: str
    titulo: str
    ancho: int
    descripcion: str


COLUMNAS_REGISTRO = [
    Columna("fecha_hora", "Fecha y hora", 20, "Momento en que terminó la petición."),
    Columna("accion", "Acción", 70, "Qué se hizo, en lenguaje llano. En los cambios indica el valor anterior y el nuevo."),
    Columna("metodo", "Método", 9, "GET = leer, PUT = modificar, POST = crear o iniciar sesión."),
    Columna("ruta", "Ruta", 48, "Dirección de la API a la que se llamó."),
    Columna("http", "HTTP", 7, "Código de respuesta: 200 = correcto; 409 = correo ya en uso; 403 = sin permiso o sesión caducada; 5xx = fallo del servidor."),
    Columna("intentos", "Intentos", 9, "Veces que se envió la petición (más de 1 si hubo fallos temporales y se reintentó)."),
    Columna("duracion_ms", "Duración (ms)", 13, "Tiempo total de la petición, incluidos los reintentos."),
    Columna("resultado", "Resultado", 60, "OK, o el error devuelto."),
]


# --- Lectura de la entrada --------------------------------------------------------

def normalizar_columna(nombre) -> str:
    """'Correo antiguo' -> 'correo_antiguo'; 'Email-Nuevo' -> 'email_nuevo'."""
    sin_tildes = unicodedata.normalize("NFKD", str(nombre or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[\s\-]+", "_", sin_tildes.strip().lower())


def _texto(valor) -> str:
    """Valor de una celda como texto limpio. Evita que un número salga como '12.0'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return str(valor).strip()


def leer_excel(ruta: str, hojas_preferidas=("Usuarios", "Resultados")) -> tuple:
    """Lee un .xlsx y devuelve (filas, nombre_de_la_hoja). Termina con un mensaje si el
    fichero no existe o no es un Excel .xlsx."""
    ruta = Path(ruta)
    if not ruta.exists():
        sys.exit(f"No existe el fichero {ruta}")
    if ruta.suffix.lower() not in (".xlsx", ".xlsm"):
        sys.exit(f"La entrada debe ser un Excel .xlsx (recibido: {ruta.name})")
    libro = load_workbook(ruta, read_only=True, data_only=True)
    hoja = next((libro[h] for h in hojas_preferidas if h in libro.sheetnames), libro.worksheets[0])

    cabecera, filas = None, []
    for numero, valores in enumerate(hoja.iter_rows(values_only=True), start=1):
        textos = [_texto(v) for v in valores]
        if not any(textos):
            continue  # fila vacía
        if cabecera is None:
            cabecera = [normalizar_columna(t) for t in textos]
            continue
        fila = {col: val for col, val in zip(cabecera, textos) if col}
        fila["_fila"] = numero
        filas.append(fila)
    libro.close()
    return filas, hoja.title


# --- Informe --------------------------------------------------------------------

def formatear_duracion(segundos: float) -> str:
    """12.34 -> '12.3 s'; 125 -> '2 min 5 s'."""
    if segundos < 60:
        return f"{segundos:.1f} s"
    minutos, resto = divmod(round(segundos), 60)
    return f"{minutos} min {resto} s"


def seccion_peticiones(registro: list) -> tuple:
    """Sección "Peticiones a la API" del resumen, calculada a partir del registro."""
    return ("Peticiones a la API", [
        ("Peticiones realizadas", len(registro)),
        ("Peticiones con error", sum(1 for r in registro if str(r["resultado"]).startswith("ERROR"))),
        ("Reintentos por fallos temporales", sum(max(int(r["intentos"] or 1) - 1, 0) for r in registro)),
        ("Nota", "Algunos errores son normales: por ejemplo, un 409 al intentar usar un correo que ya existe, "
                 "que el programa detecta y resuelve. El detalle está en la hoja Registro API."),
    ])


def _estilo_cabecera(celda) -> None:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=AZUL)
    celda.alignment = Alignment(vertical="center", wrap_text=True)


def _hoja_tabla(libro: Workbook, titulo: str, columnas: list, filas: list, colorear_estado: bool) -> None:
    """Hoja con una tabla: cabecera azul fija, filtros y anchos de columna."""
    hoja = libro.create_sheet(titulo)
    for i, columna in enumerate(columnas, start=1):
        _estilo_cabecera(hoja.cell(row=1, column=i, value=columna.titulo))
        hoja.column_dimensions[get_column_letter(i)].width = columna.ancho
    for r, fila in enumerate(filas, start=2):
        relleno = None
        if colorear_estado and fila.get("estado") in COLORES_ESTADO:
            relleno = PatternFill("solid", fgColor=COLORES_ESTADO[fila["estado"]])
        for c, columna in enumerate(columnas, start=1):
            celda = hoja.cell(row=r, column=c, value=fila.get(columna.clave, ""))
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            if relleno:
                celda.fill = relleno
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{max(len(filas) + 1, 2)}"


def _hoja_secciones(libro: Workbook, titulo: str, encabezado: str, secciones: list,
                    titulos_columnas=("Concepto", "Valor")) -> None:
    """Hoja de dos columnas organizada en secciones: [(titulo_seccion, [(concepto, valor), ...]), ...]."""
    hoja = libro.create_sheet(titulo)
    hoja.column_dimensions["A"].width = 42
    hoja.column_dimensions["B"].width = 100
    hoja.cell(row=1, column=1, value=encabezado).font = Font(bold=True, size=14, color=AZUL)
    fila = 3
    for titulo_seccion, pares in secciones:
        for col, texto in enumerate((titulo_seccion, ""), start=1):
            _estilo_cabecera(hoja.cell(row=fila, column=col, value=texto))
        fila += 1
        for concepto, valor in pares:
            celda_concepto = hoja.cell(row=fila, column=1, value=concepto)
            celda_concepto.font = Font(bold=True)
            celda_concepto.alignment = Alignment(vertical="top", wrap_text=True)
            if concepto in COLORES_ESTADO:  # los estados llevan su color también aquí
                celda_concepto.fill = PatternFill("solid", fgColor=COLORES_ESTADO[concepto])
            # Alineado a la izquierda: si no, los números quedan al fondo de la columna ancha.
            hoja.cell(row=fila, column=2, value=valor).alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True)
            fila += 1
        fila += 1


def guardar_informe(ruta: Path, titulo: str, resumen: list, columnas: list, filas: list,
                    registro: list, leyenda: list) -> Path:
    """Escribe el informe Excel completo y devuelve su ruta.

    Parámetros:
        ruta:     fichero .xlsx a crear.
        titulo:   título que aparece arriba en la hoja Resumen.
        resumen:  secciones de la hoja Resumen: [(titulo, [(concepto, valor), ...]), ...].
        columnas: lista de Columna con las columnas de la hoja Resultados.
        filas:    una dict por fila de resultados (claves = Columna.clave).
        registro: entradas del registro de peticiones (ClienteEDP.registro).
        leyenda:  secciones propias de la hoja Leyenda (estados, casos...). Las
                  descripciones de las columnas se añaden solas.
    """
    libro = Workbook()
    libro.remove(libro.active)
    _hoja_secciones(libro, "Resumen", titulo, resumen)
    _hoja_tabla(libro, "Resultados", columnas, filas, colorear_estado=True)
    _hoja_tabla(libro, "Registro API", COLUMNAS_REGISTRO, registro, colorear_estado=False)
    _hoja_secciones(libro, "Leyenda", "Cómo leer este informe", leyenda + [
        ("Columnas de la hoja Resultados", [(c.titulo, c.descripcion) for c in columnas]),
        ("Columnas de la hoja Registro API", [(c.titulo, c.descripcion) for c in COLUMNAS_REGISTRO]),
    ])
    libro.save(ruta)
    return ruta


# --- Plantilla de entrada -----------------------------------------------------------

def crear_plantilla(ruta: Path) -> Path:
    """Crea el Excel de entrada vacío para cambiar_correos.py."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Usuarios"
    for i, titulo in enumerate(COLUMNAS_PLANTILLA, start=1):
        _estilo_cabecera(hoja.cell(row=1, column=i, value=titulo))
        hoja.column_dimensions[get_column_letter(i)].width = 38 if i <= 2 else 22
    hoja.freeze_panes = "A2"
    # Desplegable con los roles admitidos en la columna Rol (E).
    validacion = DataValidation(type="list", formula1='"' + ",".join(ROLES_PLANTILLA) + '"', allow_blank=True)
    validacion.error = "Elige alumno, profesor o admin colegio"
    hoja.add_data_validation(validacion)
    validacion.add("E2:E5000")

    instrucciones = libro.create_sheet("Instrucciones")
    instrucciones.column_dimensions["A"].width = 120
    lineas = [
        ("Cómo rellenar la hoja Usuarios", True),
        ("Una fila por usuario. No cambies los títulos de la primera fila.", False),
        ("", False),
        ("Correo antiguo: correo de acceso actual del usuario. Si no tiene (p. ej. alguien que se incorpora),", False),
        ("    déjalo vacío: se tratará como un usuario nuevo (caso 3, o caso 4 si el correo nuevo ya existe).", False),
        ("Correo nuevo (obligatorio): correo de acceso que debe tener.", False),
        ("Nombre, Apellidos y Rol (opcionales): solo se copian al informe cuando no existe ninguno de los", False),
        ("    dos correos (caso 3), para saber a quién hay que avisar. El programa nunca crea usuarios.", False),
        ("Rol: alumno, profesor o admin colegio (hay un desplegable).", False),
        ("", False),
        ("Qué hace el programa con cada fila", True),
        ("Caso 1. El antiguo existe y el nuevo está libre: se pone el correo nuevo al usuario.", False),
        ("Caso 2. El antiguo existe y el nuevo ya está registrado: si el usuario que tiene el correo nuevo", False),
        ("    no tiene licencias activas, pasa a <correo nuevo>_old y después se aplica el caso 1.", False),
        ("    Si tiene alguna licencia activa, no se toca a ninguno de los dos usuarios.", False),
        ("Caso 3. El antiguo no existe y el nuevo está libre: no se crea el usuario. Queda PENDIENTE para", False),
        ("    avisar a quien corresponda y que se registre o se cargue aparte con sus licencias.", False),
        ("Caso 4. El antiguo no existe y el nuevo ya está registrado: no se hace nada.", False),
        ("", False),
        ("Ejemplo", True),
        ("Correo antiguo: ana.garcia@colegio-antiguo.es   Correo nuevo: ana.garcia@colegio-nuevo.es", False),
    ]
    for i, (texto, negrita) in enumerate(lineas, start=1):
        instrucciones.cell(row=i, column=1, value=texto).font = Font(bold=negrita, size=12 if negrita else 11)
    libro.save(ruta)
    return ruta


if __name__ == "__main__":
    destino = Path(sys.argv[1] if len(sys.argv) > 1 else "plantilla_entrada.xlsx")
    print("Plantilla creada:", crear_plantilla(destino))
